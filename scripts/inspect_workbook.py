from __future__ import annotations

import json
import sys
from pathlib import Path

from openpyxl import load_workbook


def serialise(value):
    if value is None:
        return None
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


source = Path(sys.argv[1])
sys.stdout.reconfigure(encoding="utf-8")
workbook = load_workbook(source, data_only=False)

report = {
    "file": str(source),
    "worksheets": [],
    "defined_names": [str(name) for name in workbook.defined_names.values()],
}

for sheet in workbook.worksheets:
    cells = []
    for row in sheet.iter_rows():
        for cell in row:
            if cell.value is None and cell.comment is None and cell.hyperlink is None:
                continue
            cells.append(
                {
                    "coordinate": cell.coordinate,
                    "value": serialise(cell.value),
                    "data_type": cell.data_type,
                    "number_format": cell.number_format,
                    "comment": cell.comment.text if cell.comment else None,
                    "hyperlink": cell.hyperlink.target if cell.hyperlink else None,
                    "style_id": cell.style_id,
                }
            )

    validations = []
    for validation in sheet.data_validations.dataValidation:
        validations.append(
            {
                "type": validation.type,
                "formula1": validation.formula1,
                "formula2": validation.formula2,
                "allow_blank": validation.allow_blank,
                "sqref": str(validation.sqref),
                "prompt": validation.prompt,
                "error": validation.error,
            }
        )

    report["worksheets"].append(
        {
            "title": sheet.title,
            "state": sheet.sheet_state,
            "dimensions": sheet.calculate_dimension(),
            "freeze_panes": str(sheet.freeze_panes) if sheet.freeze_panes else None,
            "merged_cells": [str(value) for value in sheet.merged_cells.ranges],
            "validations": validations,
            "tables": list(sheet.tables.keys()),
            "images": len(sheet._images),
            "cells": cells,
        }
    )

print(json.dumps(report, indent=2, ensure_ascii=False))
